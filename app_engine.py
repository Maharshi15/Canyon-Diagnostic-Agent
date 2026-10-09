"""Canyon Diagnostic Agent: Application Diagnostic engine.

The vision model reviews screenshots (Evidence level 1, Observed).
Analytics files are analysed by code here (Evidence level 2, Measured).
"""
import base64
import io
import json
import re

import pandas as pd

MAX_SCREENS = 12
SEV_RANK = {"High": 0, "Medium": 1, "Low": 2}

SYSTEM = """You are the Canyon Diagnostic Agent of Canyon Data Labs, reviewing a business application from screenshots.
The screenshots are given in the order of the user journey. Review them as a senior UX (User Experience) and process consultant.

Rules:
* Report only what is visible in the screenshots. Every finding must name the screen and quote or describe the visible evidence.
* Screenshots are Evidence level 1 (Observed). Never claim a page is unused, never invent drop off rates, usage numbers or business impact figures.
* Look for: steps that can merge, fields asked twice, fields that can be derived or auto filled from data already entered
  (for example city and state from PIN code, bank and branch from IFSC code), unnecessary uploads, missing save or draft,
  unclear labels, low contrast text, small tap targets, inconsistent navigation, cluttered screens, long forms, missing help or validation hints.
* Propose a simplified journey with fewer steps where it is realistic, and say what is removed, merged or auto filled.
* For each important simplification, propose an A/B experiment to validate it.
* Plain, simple English. Never use hyphen or dash characters in your text. Spell out each abbreviation in brackets the first time.
* Severity: High = likely to block or lose users, or cause wrong data. Medium = slows users or adds rework. Low = polish.
* Effort: S, M or L.

Return only valid JSON with exactly this structure:
{
 "app_summary": "two or three sentences on what the app does in this journey and the overall impression",
 "screens": [{"index": 1, "name": "", "purpose": "", "fields": 0, "required_fields": 0, "uploads": 0, "notes": ""}],
 "findings": [{"screen": 1, "title": "", "category": "Form|Navigation|Content|Visual|Consistency|Accessibility|Mobile|Process",
               "severity": "High|Medium|Low", "evidence": "", "recommendation": "", "effort": "S|M|L"}],
 "journey": {"current_steps": 0, "proposed_steps": 0,
             "proposed_flow": [{"step": 1, "name": "", "what_happens": ""}],
             "changes": ["what is removed, merged or auto filled"]},
 "experiments": [{"hypothesis": "", "variant": "", "primary_metric": "", "guardrail_metric": "",
                  "platforms": "for example Optimizely, VWO, Statsig, GrowthBook or Firebase A/B Testing"}],
 "needs_more_evidence": ["what analytics, usage logs or tests would confirm the findings"]
}"""


def prepare_images(files, max_side=1280):
    """files: list of (name, bytes). Returns list of dicts with name and base64 JPEG."""
    from PIL import Image
    out = []
    for name, data in files[:MAX_SCREENS]:
        im = Image.open(io.BytesIO(data))
        im = im.convert("RGB")
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=85)
        out.append({"name": name, "b64": base64.b64encode(buf.getvalue()).decode()})
    return out


def user_text(app_name, app_type, journey, users, extra):
    return (f"Application: {app_name or 'Not specified'}\nType: {app_type}\n"
            f"Journey shown: {journey or 'Not specified'}\nMain users: {users or 'Not specified'}\n"
            f"Context from Canyon: {extra or 'None'}\n"
            "Screenshots follow in journey order, numbered from 1. Return the JSON only.")


def _azure(cfg, text, images):
    from openai import AzureOpenAI, OpenAI, BadRequestError
    endpoint = cfg["endpoint"]
    for suffix in ("/openai/v1", "/openai"):
        if endpoint.endswith(suffix):
            endpoint = endpoint[: -len(suffix)]
    if cfg["api_version"]:
        client = AzureOpenAI(azure_endpoint=endpoint, api_key=cfg["key"], api_version=cfg["api_version"])
    else:
        client = OpenAI(base_url=endpoint + "/openai/v1/", api_key=cfg["key"])
    content = [{"type": "text", "text": text}]
    for i, im in enumerate(images, 1):
        content.append({"type": "text", "text": f"Screen {i}: {im['name']}"})
        content.append({"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + im["b64"]}})
    messages = [{"role": "system", "content": SYSTEM}, {"role": "user", "content": content}]
    kwargs = {"model": cfg["deployment"], "messages": messages, "max_completion_tokens": 12000,
              "response_format": {"type": "json_object"}}
    for _ in range(3):
        try:
            resp = client.chat.completions.create(**kwargs)
            return resp.choices[0].message.content or ""
        except BadRequestError as e:
            msg = str(e)
            if "response_format" in msg and "response_format" in kwargs:
                kwargs.pop("response_format")
            elif "max_completion_tokens" in msg and "max_completion_tokens" in kwargs:
                kwargs.pop("max_completion_tokens")
                kwargs["max_tokens"] = 4000
            else:
                raise
    raise RuntimeError("The AI service rejected the request settings.")


def _anthropic(cfg, text, images):
    import anthropic
    content = [{"type": "text", "text": text}]
    for i, im in enumerate(images, 1):
        content.append({"type": "text", "text": f"Screen {i}: {im['name']}"})
        content.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": im["b64"]}})
    msg = anthropic.Anthropic(api_key=cfg["key"]).messages.create(
        model=cfg["model"], max_tokens=8000, system=SYSTEM, messages=[{"role": "user", "content": content}])
    return "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")


def parse_json(text):
    text = (text or "").strip()
    try:
        return json.loads(text)
    except Exception:
        pass
    m = re.search(r"\{.*\}", text, re.S)
    if not m:
        raise ValueError("The AI did not return a readable result. Try again.")
    return json.loads(m.group(0))


def normalize(result, n_screens):
    """Make the AI result safe to display: fill gaps, clamp values, sort findings."""
    r = dict(result or {})
    r.setdefault("app_summary", "")
    screens = [s for s in (r.get("screens") or []) if isinstance(s, dict)]
    r["screens"] = screens
    fixed = []
    for f in r.get("findings") or []:
        if not isinstance(f, dict):
            continue
        sev = str(f.get("severity", "Medium")).title()
        f["severity"] = sev if sev in SEV_RANK else "Medium"
        try:
            f["screen"] = max(1, min(int(f.get("screen", 1)), n_screens or 1))
        except Exception:
            f["screen"] = 1
        f["effort"] = str(f.get("effort", "M")).upper()[:1] or "M"
        for k in ("title", "category", "evidence", "recommendation"):
            f[k] = str(f.get(k, "") or "")
        f["evidence_level"] = "1 Observed"
        f["measured"] = ""
        fixed.append(f)
    fixed.sort(key=lambda f: (SEV_RANK[f["severity"]], f["screen"]))
    r["findings"] = fixed
    j = r.get("journey") or {}
    j.setdefault("current_steps", n_screens)
    j.setdefault("proposed_steps", j.get("current_steps"))
    j.setdefault("proposed_flow", [])
    j.setdefault("changes", [])
    r["journey"] = j
    r["experiments"] = [e for e in (r.get("experiments") or []) if isinstance(e, dict)]
    r["needs_more_evidence"] = [str(x) for x in (r.get("needs_more_evidence") or [])]
    return r


def analyse_screens(cfg, files, app_name="", app_type="", journey="", users="", extra=""):
    images = prepare_images(files)
    text = user_text(app_name, app_type, journey, users, extra)
    raw = _azure(cfg, text, images) if cfg["provider"] == "azure" else _anthropic(cfg, text, images)
    return normalize(parse_json(raw), len(images))


# Analytics (Evidence level 2, measured by code)

def _num(series):
    return pd.to_numeric(series.astype(str).str.replace(",", "").str.strip(), errors="coerce")


def funnel_analysis(df):
    """df: first column step name, second column users who completed it.
    Row 1 is users who started the journey; row k+1 is users who completed screen k."""
    if df is None or df.shape[1] < 2 or len(df) < 2:
        return None
    steps = df.iloc[:, 0].astype(str).tolist()
    users = _num(df.iloc[:, 1]).fillna(0).tolist()
    rows = []
    for i in range(1, len(steps)):
        prev, cur = users[i - 1], users[i]
        drop = (prev - cur) / prev * 100 if prev else 0
        rows.append({"Screen": i, "Step": steps[i], "Users completing": int(cur),
                     "Lost at this step": int(max(prev - cur, 0)), "Drop off %": round(drop, 1)})
    start, end = users[0], users[-1]
    worst = max(rows, key=lambda r: r["Drop off %"])
    return {"rows": rows, "start": int(start), "end": int(end),
            "completion": round(end / start * 100, 1) if start else 0.0, "worst": worst}


def page_usage_analysis(df, threshold_pct=0.5):
    if df is None or df.shape[1] < 2 or len(df) < 2:
        return None
    pages = df.iloc[:, 0].astype(str).tolist()
    views = _num(df.iloc[:, 1]).fillna(0).tolist()
    total = sum(views) or 1
    rows = []
    for p, v in sorted(zip(pages, views), key=lambda x: -x[1]):
        share = v / total * 100
        rows.append({"Page": p, "Views": int(v), "Share of all views %": round(share, 2),
                     "Signal": "Very low usage: review, merge or retire" if share < threshold_pct else ""})
    return {"rows": rows, "total": int(total), "low": [r for r in rows if r["Signal"]], "threshold": threshold_pct}


def attach_measured(result, funnel, drop_threshold=15.0):
    """Raise findings to Evidence level 2 where the funnel shows heavy drop off on the same screen."""
    if not funnel:
        return result
    by_screen = {r["Screen"]: r for r in funnel["rows"]}
    for f in result["findings"]:
        row = by_screen.get(f["screen"])
        if row and row["Drop off %"] >= drop_threshold:
            f["evidence_level"] = "2 Measured"
            f["measured"] = f"{row['Drop off %']}% of users who started this step did not complete it ({row['Lost at this step']} users)"
    return result


def counts(result):
    c = {"High": 0, "Medium": 0, "Low": 0}
    for f in result["findings"]:
        c[f["severity"]] += 1
    return c


def to_excel(result, funnel, usage, buf, client="", app_name=""):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    c = counts(result)
    j = result["journey"]
    summary = [["Canyon Diagnostic Agent: Application Diagnostic", ""],
               ["Client", client or "Not specified"], ["Application", app_name or "Not specified"],
               ["Screens reviewed", len(result["screens"]) or j.get("current_steps")],
               ["Findings (High / Medium / Low)", f"{c['High']} / {c['Medium']} / {c['Low']}"],
               ["Journey steps: current to proposed", f"{j.get('current_steps')} to {j.get('proposed_steps')}"],
               ["Overview", result.get("app_summary", "")]]
    if funnel:
        summary += [["Measured completion rate", f"{funnel['completion']}% ({funnel['end']} of {funnel['start']})"],
                    ["Largest drop off", f"{funnel['worst']['Step']}: {funnel['worst']['Drop off %']}%"]]
    summary += [["", ""], ["Evidence levels", "1 Observed (screenshots), 2 Measured (analytics), 3 Validated (experiment)"],
                ["Note", "Screenshot findings suggest improvements. Confirm with analytics and an A/B test before rollout."]]
    sheets = {
        "Summary": pd.DataFrame(summary, columns=["Item", "Value"]),
        "Findings": pd.DataFrame([{"Severity": f["severity"], "Screen": f["screen"], "Finding": f["title"],
                                   "Category": f["category"], "Evidence level": f["evidence_level"],
                                   "Evidence": f["evidence"] + (f" Measured: {f['measured']}." if f["measured"] else ""),
                                   "Recommendation": f["recommendation"], "Effort": f["effort"]}
                                  for f in result["findings"]]),
        "Screens": pd.DataFrame([{"Screen": s.get("index"), "Name": s.get("name"), "Purpose": s.get("purpose"),
                                  "Fields": s.get("fields"), "Required": s.get("required_fields"),
                                  "Uploads": s.get("uploads"), "Notes": s.get("notes")} for s in result["screens"]]),
        "Proposed Journey": pd.DataFrame([{"Step": p.get("step"), "Name": p.get("name"), "What happens": p.get("what_happens")}
                                          for p in j.get("proposed_flow", [])] +
                                         [{"Step": "Change", "Name": "", "What happens": ch} for ch in j.get("changes", [])]),
        "Experiments": pd.DataFrame([{"Hypothesis": e.get("hypothesis"), "Variant": e.get("variant"),
                                      "Primary metric": e.get("primary_metric"), "Guardrail metric": e.get("guardrail_metric"),
                                      "Platforms": e.get("platforms")} for e in result["experiments"]]),
    }
    if funnel:
        sheets["Funnel (Measured)"] = pd.DataFrame(funnel["rows"])
    if usage:
        sheets["Page Usage (Measured)"] = pd.DataFrame(usage["rows"])
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            (df if len(df) else pd.DataFrame({"Info": ["Nothing reported"]})).to_excel(xw, sheet_name=name, index=False)
        head = PatternFill("solid", fgColor="0B6E69")
        fills = {"High": "F8E1DE", "Medium": "F6EAD3", "Low": "E0E9F3"}
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head
            ws.freeze_panes = "A2"
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(x.value or "")) for x in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 60)
                for x in col:
                    x.alignment = Alignment(wrap_text=True, vertical="top")
        for row in xw.book["Findings"].iter_rows(min_row=2):
            if row[0].value in fills:
                row[0].fill = PatternFill("solid", fgColor=fills[row[0].value])
