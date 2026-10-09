"""Pilot Request Kit: the message asking a pilot client for data, and the data request checklist.

House style: plain English, no dash characters, abbreviations spelled out the first time in each output.
"""
import io
from datetime import date

import streamlit as st

from report_engine import TEAL, clean

SIGNATURE = ("Best regards,\nMaharshi Upadhyay\nBusiness Development Head\nCanyon Data Labs\n"
             "maharshi@canyondatalabs.com")

FOCUS = {"Master Data": "master data", "Application": "business apps", "Process": "workflows"}
ASK_SHORT = {
    "Master Data": "a sample extract of one master in Excel with personal data masked",
    "Application": "screenshots of one important journey in the app with personal data blurred",
    "Process": "the SOP (Standard Operating Procedure) or a 30 minute walkthrough of one process",
}
CHECKLIST = {
    "Master Data": [
        "An extract of the material, customer or distributor master as Excel or CSV (Comma Separated Values). "
        "500 to 5,000 records is enough to start.",
        "Columns: code, description, unit of measure, group or category, plant or branch. For customer or distributor "
        "masters also GSTIN (Goods and Services Tax Identification Number), PIN (Postal Index Number) code and city.",
        "The approved list of units of measure and groups, if one exists.",
        "Personal data: leave out phone numbers and email addresses, or mask them.",
    ],
    "Application": [
        "Screenshots of one important journey, in order, for example registering a retailer or placing an order. "
        "Name them 01, 02, 03 so the order is kept. Up to 12 screens.",
        "Blur customer names, phone numbers, account numbers and photos of documents before sharing.",
        "Optional, makes findings measured instead of observed: a funnel export (users who started the journey, then "
        "users who completed each screen) and page views for the last 90 days, from GA4 (Google Analytics 4), "
        "Mixpanel, Firebase or the app's own logs.",
        "The names of any analytics or A/B (split) testing tools you already use.",
    ],
    "Process": [
        "The SOP (Standard Operating Procedure), a flowchart, or a 30 minute walkthrough with the process owner.",
        "For each step: who does it, in which system or tool, roughly how long the work takes and how long it waits.",
        "The approximate number of cases per month.",
        "Optional, makes findings measured instead of estimated: timestamps for one month of cases from the workflow "
        "tool or ERP (Enterprise Resource Planning) system.",
    ],
}
GENERAL = [
    "One contact person who can answer questions during the pilot.",
    "A 30 minute call to review the findings with your team.",
    "Share files by email or through a shared folder of your choice. A sample is enough, the full dataset is not needed.",
]
WHAT_YOU_GET = [
    "A Diagnostic Report with an evidence level on every finding, quick wins for the next 30 days and a short roadmap.",
    "Excel workbooks with every finding and the records or screens behind it.",
    "A review call with Canyon. Nothing in your systems is changed during the pilot.",
]
DATA_LINE = ("We use the files only for this diagnostic, personal data can be masked or left out, and we delete the "
             "files when the pilot ends. We are happy to sign your NDA (Non Disclosure Agreement) first.")
CONSENT = ("We confirm that Canyon Data Labs may use the files we share for this diagnostic pilot only. "
           "Personal data in the files is masked or left out. Canyon Data Labs will delete the files when the pilot ends.")


def _join(items):
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def messages(first, company, modules, detail="", free=False):
    first, company = first or "[first name]", company or "[Company]"
    focus = _join([FOCUS[m] for m in modules])
    asks = _join([ASK_SHORT[m] for m in modules])
    pilot = "a pilot diagnostic at no cost" if free else "a pilot diagnostic"
    d = detail.strip().rstrip(".")
    offer = (f"{d[0].upper() + d[1:]}, which is why I thought of {company} for one of the few pilots we are running "
             f"this quarter. I would like to offer you {pilot}." if d else
             f"We are running a few pilots this quarter and I would like to offer {pilot} to {company}.")

    linkedin = (f"Hi {first},\n\n"
                "I am Maharshi, Business Development Head at Canyon Data Labs. We are a data engineering and AI "
                "(Artificial Intelligence) company based out of Ahmedabad, India.\n\n"
                f"We have built a diagnostic that reviews a company's {focus} and shows what to fix, what to simplify "
                "and where automation can save time.\n\n"
                f"{offer} All we would need is {asks}. You get a short report with the findings and quick wins, and "
                "nothing in your systems changes.\n\n"
                "Would you be open to a short call to explore this?")

    bullets = "\n".join(f"* {ASK_SHORT[m][0].upper() + ASK_SHORT[m][1:]}" for m in modules)
    subject = f"A pilot diagnostic of {company}'s {focus}"
    email = (f"Subject: {subject}\n\nHi {first},\n\n"
             "I am Maharshi, Business Development Head at Canyon Data Labs. We are a data engineering and AI "
             "(Artificial Intelligence) company based out of Ahmedabad, India.\n\n"
             f"We have built the Canyon Diagnostic Agent. It reviews a company's {focus} and shows what to fix, what to "
             "simplify and where automation or AI can create measurable value.\n\n"
             f"{offer} To start, we would need:\n\n"
             f"{bullets}\n\n"
             "You receive a Diagnostic Report with an evidence level on every finding, quick wins for the next 30 days "
             "and a short roadmap. Canyon works with the systems you already run on, no need to replace anything.\n\n"
             f"{DATA_LINE}\n\n"
             "Would you be open to a 20 minute call next week to agree the scope? I can also share the full data "
             "checklist in advance.\n\n"
             f"{SIGNATURE}")

    follow = (f"Hi {first},\n\n"
              f"Following up on my note about a pilot diagnostic for {company}. It needs very little from your team, "
              f"just {asks}, and you get a short report with the quick wins.\n\n"
              "Would a 20 minute call next week work to agree the scope?")
    return {"LinkedIn message": linkedin, "Email": email, "Follow up after five days": follow}


def checklist_docx(company, modules, client_contact=""):
    import docx
    from docx.shared import Pt, RGBColor

    doc = docx.Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string(TEAL)

    def p(text, bold=False, size=None):
        r = doc.add_paragraph().add_run(clean(text))
        r.bold = bold
        if size:
            r.font.size = Pt(size)

    def items(lst):
        for x in lst:
            doc.add_paragraph(clean(x), style="List Bullet")

    p("CANYON DATA LABS", bold=True, size=9)
    doc.add_heading(f"Pilot diagnostic: data request for {company or 'your company'}", level=1)
    p(f"Date: {date.today():%d %B %Y}")
    p("Thank you for agreeing to a pilot. This page lists what we need, how to share it, and what you get back.")
    for m in modules:
        doc.add_heading(f"{m}", level=2)
        items(CHECKLIST[m])
    doc.add_heading("For every pilot", level=2)
    items(GENERAL)
    doc.add_heading("How we handle your data", level=2)
    items([DATA_LINE, "Screenshots and documents are reviewed with the help of an enterprise AI (Artificial Intelligence) "
                      "service. Only the files you share are used, and only for this diagnostic."])
    doc.add_heading("What you get", level=2)
    items(WHAT_YOU_GET)
    doc.add_heading("Pilot confirmation", level=2)
    p("Please reply to the email with this text, or sign below, before sharing files.")
    p(CONSENT)
    for label in ("Name", "Role", "Company", "Date", "Signature"):
        p(f"{label}: " + ("_" * 40 if label != "Company" or not company else company))
    p("Contact at Canyon: Maharshi Upadhyay, Business Development Head, maharshi@canyondatalabs.com")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render(client):
    st.subheader("Pilot Request Kit")
    st.write("Drafts the message asking a pilot client for data, and the data request checklist to send with it.")
    c1, c2 = st.columns(2)
    first = c1.text_input("Contact first name", placeholder="For example: Rohan")
    company = c2.text_input("Company", value=client or "", placeholder="For example: ABC Manufacturing")
    modules = st.multiselect("What the pilot covers", list(FOCUS), default=["Master Data"])
    detail = st.text_input("One real, specific detail about them (optional, makes the message personal)",
                           placeholder="For example: I saw that ABC recently added a second plant in Sanand")
    free = st.checkbox("Say the pilot is at no cost", value=False,
                       help="Tick only if Canyon has agreed to run this pilot free of charge.")
    if not modules:
        st.info("Choose at least one module.")
        return

    st.caption("Before sending: check the specific detail is true, and confirm the data handling lines with the Canyon "
               "Tech Lead until the security review is done.")
    for title, text in messages(first, company, modules, detail, free).items():
        st.markdown(f"**{title}**")
        st.code(clean(text), language=None, wrap_lines=True)

    st.markdown("#### Data request checklist")
    for m in modules:
        st.markdown(f"**{m}**")
        for x in CHECKLIST[m]:
            st.markdown(f"* {x}")
    st.markdown("**For every pilot**")
    for x in GENERAL:
        st.markdown(f"* {x}")
    st.download_button("Download checklist (Word)", checklist_docx(company, modules), type="primary",
                       file_name=f"Canyon_Pilot_Data_Request_{(company or 'Client').replace(' ', '_')}.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
