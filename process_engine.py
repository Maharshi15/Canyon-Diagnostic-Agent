"""Canyon Diagnostic Agent: Process Diagnostic engine.

A process is a table of steps: step, owner, system, type, work minutes, wait hours, share of cases.
All counts, times and ratings come from this file, so results are repeatable.
The AI (Artificial Intelligence) model is used only to draft the step table from an SOP (Standard Operating Procedure).
"""
import io
import json
import re

import pandas as pd

SEV_RANK = {"High": 0, "Medium": 1, "Low": 2}
TYPES = ["Task", "Data entry", "Check", "Approval", "Handoff"]
WORKAROUND = re.compile(r"excel|spreadsheet|sheet|e ?mail|outlook|gmail|whatsapp|paper|phone|call|sms|register|manual", re.I)
INTAKE = re.compile(r"whatsapp|paper|phone|call|sms", re.I)
APPROACH = {
    "Data entry": ("Enter once: integration or AI (Artificial Intelligence) document capture fills this system", "M"),
    "Check": ("Automated validation rules, only exceptions go to a person", "S"),
    "Approval": ("Rules based auto approval within agreed limits, one click mobile approval for the rest", "M"),
    "Handoff": ("Automatic notification and status update, no manual follow up", "S"),
    "Task": ("Workflow automation or an AI agent takes over the routine part", "M"),
}
TIME_SOURCES = ["Estimates from interviews or the SOP (Evidence level 1, Observed)",
                "System timestamps or logs (Evidence level 2, Measured)"]

DRAFT_SYSTEM = """You are the Canyon Diagnostic Agent of Canyon Data Labs. You turn a process document into a step table.
Rules:
* Use only what the document says. Never invent times. If the document does not give a time, use 0.
* One row per step, in the order the work happens. Name the owner (role, not a person's name) and the system or tool used.
* type must be one of: Task, Data entry, Check, Approval, Handoff.
* Plain, simple English. Never use hyphen or dash characters.
Return only valid JSON with exactly this structure:
{"process_name": "", "steps": [{"step": "", "owner": "", "system": "", "type": "Task", "work_minutes": 0, "wait_hours": 0}],
 "assumptions": ["anything unclear in the document that the client should confirm"]}"""


# Reading inputs

def read_document(name, data):
    """Plain text from an SOP file: .txt, .md, .docx or .pdf."""
    name = name.lower()
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    return data.decode("utf-8", errors="ignore")


def draft_steps(cfg, text, process_name=""):
    """Ask the AI to turn SOP text into a step table. Returns (DataFrame, assumptions)."""
    from ai_summary import ask
    prompt = f"Process: {process_name or 'Not specified'}\n\nDocument:\n{text[:30000]}\n\nReturn the JSON only."
    raw = ask(cfg, DRAFT_SYSTEM, prompt, max_tokens=4000, json_mode=True)
    try:
        out = json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw or "", re.S)
        if not m:
            raise ValueError("The AI did not return a readable step table. Try again.")
        out = json.loads(m.group(0))
    rows = [{"Step": s.get("step", ""), "Owner": s.get("owner", ""), "System": s.get("system", ""),
             "Type": s.get("type", "Task"), "Work minutes": s.get("work_minutes", 0),
             "Wait hours": s.get("wait_hours", 0), "Share of cases %": 100}
            for s in out.get("steps") or [] if isinstance(s, dict)]
    return pd.DataFrame(rows), [str(a) for a in out.get("assumptions") or []]


def template():
    return pd.DataFrame([{"Step": "", "Owner": "", "System": "", "Type": "Task", "Work minutes": 0,
                          "Wait hours": 0, "Share of cases %": 100}])


def _map_columns(columns):
    cols = [str(c) for c in columns]
    taken = set()

    def find(pattern):
        for i, c in enumerate(cols):
            if i not in taken and re.search(pattern, c.strip(), re.I):
                taken.add(i)
                return i
        return -1

    m = {"step": find(r"step|activity|task|description"), "work": find(r"work|effort|touch|handling|minute"),
         "wait": find(r"wait|delay|queue|idle|hour"), "share": find(r"share|percent|%"),
         "owner": find(r"owner|role|who|team|department"), "system": find(r"system|tool|application|app"),
         "type": find(r"type|kind"), "manual": find(r"manual|automat")}
    if m["step"] < 0:
        m["step"] = 0
    return m


def _type(value, step):
    for text in (str(value or ""), str(step or "")):
        t = text.lower()
        if t.strip().title() in TYPES or t.strip() == "data entry":
            return "Data entry" if t.strip() == "data entry" else t.strip().title()
        if re.search(r"approv|sanction|sign off", t):
            return "Approval"
        if re.search(r"entry|enter|key in|input|record in|update in|post", t):
            return "Data entry"
        if re.search(r"check|verif|valid|review|match|reconcil", t):
            return "Check"
        if re.search(r"hand ?over|send|inform|notify|forward|share|email|mail", t):
            return "Handoff"
    return "Task"


def _num(v):
    try:
        return max(float(str(v).replace(",", "").strip() or 0), 0.0)
    except ValueError:
        return 0.0


def _manual(v):
    t = str(v or "").strip().lower()
    return not re.match(r"^(no|n|false|0|auto|automated|automatic|system)$", t)


def clean_steps(df):
    """Map any reasonable column names to the standard step table."""
    m = _map_columns(df.columns)
    out = []
    for _, r in df.iterrows():
        step = str(r.iloc[m["step"]]).strip()
        if not step or step.lower() in ("nan", "none"):
            continue
        get = lambda k, d="": r.iloc[m[k]] if m[k] >= 0 else d
        share = _num(get("share", 100)) if m["share"] >= 0 else 100.0
        out.append({"step": step, "owner": str(get("owner")).strip(), "system": str(get("system")).strip(),
                    "type": _type(get("type"), step), "work_min": _num(get("work", 0)), "wait_h": _num(get("wait", 0)),
                    "share": min(share, 100.0) if share else 100.0, "manual": _manual(get("manual", "yes"))})
    for i, s in enumerate(out, 1):
        s["n"] = i
        s["owner"] = "" if s["owner"].lower() in ("nan", "none") else s["owner"]
        s["system"] = "" if s["system"].lower() in ("nan", "none") else s["system"]
    return out


def standard_table(df):
    """Any step table in, the standard editable columns out."""
    rows = [{"Step": s["step"], "Owner": s["owner"], "System": s["system"], "Type": s["type"],
             "Work minutes": s["work_min"], "Wait hours": s["wait_h"], "Share of cases %": s["share"]}
            for s in clean_steps(df)]
    return pd.DataFrame(rows) if rows else template()


# Analysis

def analyse(df, cases_per_month=0, time_source=TIME_SOURCES[0]):
    steps = clean_steps(df)
    if not steps:
        return {"error": "The step table is empty. Add at least one step."}
    cases = max(int(cases_per_month or 0), 0)
    timed = any(s["work_min"] or s["wait_h"] for s in steps)
    time_level = "2 Measured" if "Measured" in time_source else "1 Observed"
    for s in steps:
        s["hours_month"] = round(s["work_min"] * s["share"] / 100 * cases / 60, 1) if s["manual"] else 0.0

    owners = [s["owner"] for s in steps if s["owner"]]
    handoffs = sum(1 for a, b in zip(steps, steps[1:]) if a["owner"] and b["owner"] and a["owner"] != b["owner"])
    work_h = sum(s["work_min"] * s["share"] / 100 for s in steps) / 60
    wait_h = sum(s["wait_h"] * s["share"] / 100 for s in steps)
    cycle_h = work_h + wait_h
    metrics = {"steps": len(steps), "owners": len(set(o.lower() for o in owners)),
               "systems": len(set(s["system"].lower() for s in steps if s["system"])), "handoffs": handoffs,
               "approvals": sum(1 for s in steps if s["type"] == "Approval"),
               "manual_steps": sum(1 for s in steps if s["manual"]), "timed": timed, "cases": cases,
               "work_h": round(work_h, 1), "wait_h": round(wait_h, 1), "cycle_h": round(cycle_h, 1),
               "cycle_days": round(cycle_h / 24, 1),
               "flow_eff": round(work_h / cycle_h * 100, 1) if cycle_h else None,
               "manual_h_month": round(sum(s["hours_month"] for s in steps), 1)}

    findings = []

    def add(title, category, severity, evidence, rec, effort, at, level="1 Observed"):
        findings.append({"title": title, "category": category, "severity": severity, "evidence": evidence,
                         "recommendation": rec, "effort": effort, "steps": at, "evidence_level": level})

    # 1. The same data typed into a second system
    entries = []
    for s in steps:
        if s["type"] != "Data entry":
            continue
        earlier = next((e for e in entries if e["system"] and s["system"] and e["system"].lower() != s["system"].lower()), None)
        again = re.search(r"re ?enter|again|copy|retype", s["step"], re.I)
        if earlier or again:
            src = f"step {earlier['n']} ({earlier['system']})" if earlier else "an earlier step"
            add(f"Data typed again at step {s['n']}: {s['step']}", "Re entry", "High",
                f"Step {s['n']} ({s['owner'] or 'owner not given'}) enters data into {s['system'] or 'a system'} "
                f"that was already captured at {src}.",
                f"Capture the data once and pass it to {s['system'] or 'the next system'} through an integration, "
                "or let an AI agent read the first record and fill it. This removes typing time and copy errors.",
                "M", [s["n"]])
        entries.append(s)

    # 2. Work outside the core systems
    wa = [s for s in steps if s["system"] and WORKAROUND.search(s["system"])]
    if wa:
        add("Work runs on spreadsheets, email or chat", "Workaround", "High" if len(wa) >= 3 else "Medium",
            f"{len(wa)} of {len(steps)} steps use {', '.join(dict.fromkeys(s['system'] for s in wa))}: "
            + "; ".join(f"step {s['n']} {s['step']}" for s in wa) + ".",
            "Move these steps into a workflow tool or the core system, so every case has one record, one owner "
            "and a status. This also makes cycle time measurable.", "M", [s["n"] for s in wa])

    # 3. Approvals
    ap = [s for s in steps if s["type"] == "Approval"]
    if len(ap) >= 2:
        quick = [s for s in ap if s["work_min"] <= 5 and s["wait_h"] >= 24]
        ap_wait = sum(s["wait_h"] for s in ap)
        ev = (f"{len(ap)} approval steps: " + "; ".join(
            f"step {s['n']} {s['owner'] or s['step']}" + (f" ({s['share']:g}% of cases)" if s["share"] < 100 else "")
            for s in ap) + ".")
        if timed:
            ev += f" Together they wait {ap_wait:g} hours."
        if quick:
            ev += (f" {len(quick)} of them take 5 minutes or less to review but wait 24 hours or more, "
                   "which suggests little decision value. Confirm with the approvers.")
        heavy = timed and cycle_h and ap_wait / cycle_h >= 0.4
        add(f"{len(ap)} approval steps in one process", "Approval", "High" if heavy else "Medium", ev,
            "Agree approval limits: auto approve low value or low risk cases by rule, run the remaining approvals "
            "in parallel, and allow one click approval on mobile.", "M", [s["n"] for s in ap],
            time_level if timed else "1 Observed")

    # 4. Waiting time
    if work_h and metrics["flow_eff"] is not None and metrics["flow_eff"] < 20:
        longest = sorted(steps, key=lambda s: -s["wait_h"])[:3]
        add("Cases spend most of their time waiting", "Waiting", "High" if metrics["flow_eff"] < 10 else "Medium",
            f"Active work per case is {metrics['work_h']} hours, total elapsed time is {metrics['cycle_h']} hours "
            f"({metrics['cycle_days']} days), so only {metrics['flow_eff']}% of the time is active work. Longest waits: "
            + "; ".join(f"step {s['n']} {s['step']} ({s['wait_h']:g} hours)" for s in longest) + ".",
            "Set a target time for each step, show each owner their queue, and send automatic reminders when a case "
            "waits too long.", "S", [s["n"] for s in longest], time_level)

    # 5. Handoffs
    if handoffs >= 4:
        add(f"The case changes hands {handoffs} times", "Handoff", "Medium",
            f"{metrics['owners']} different owners work on one case, with {handoffs} handoffs between them. "
            "Every handoff adds waiting and a chance of losing information.",
            "Let the person who starts the case complete more steps, and move routine checks into the system.",
            "M", [b["n"] for a, b in zip(steps, steps[1:]) if a["owner"] and b["owner"] and a["owner"] != b["owner"]])

    # 6. Manual checks
    ck = [s for s in steps if s["type"] == "Check" and s["manual"]]
    if ck:
        add("Manual checks that system rules could do", "Check", "Medium",
            f"{len(ck)} manual checks: " + "; ".join(f"step {s['n']} {s['step']}" for s in ck) + ".",
            "Write the check criteria as system rules, so only exceptions reach a person.", "S", [s["n"] for s in ck])

    findings.sort(key=lambda f: (SEV_RANK[f["severity"]], f["steps"][0] if f["steps"] else 0))

    # Automation opportunities
    opps = []
    total_h = metrics["manual_h_month"] or 1
    for s in steps:
        if not s["manual"]:
            continue
        approach, effort = APPROACH[s["type"]]
        if s["type"] == "Task" and s["system"] and INTAKE.search(s["system"]):
            approach, effort = "Digital intake form or WhatsApp Business bot, so the case starts as a clean record", "M"
        h = s["hours_month"]
        # Impact: large share of the manual hours, or a long wait before the case can move on
        share_h = h / total_h * 100
        impact = ("High" if (h >= 40 and share_h >= 15) or s["wait_h"] >= 48 else
                  "Medium" if (h >= 10 and share_h >= 5) or s["wait_h"] >= 24 else "Low")
        opps.append({"Step": s["n"], "Name": s["step"], "Type": s["type"], "Owner": s["owner"], "System": s["system"],
                     "Approach": approach, "Effort": effort, "Impact": impact,
                     "Manual hours per month": h, "Wait hours per case": s["wait_h"],
                     "Quick win": "Yes" if effort == "S" and impact != "Low" else ""})
    opps.sort(key=lambda o: (SEV_RANK[o["Impact"]], -o["Manual hours per month"], -o["Wait hours per case"]))

    experiments = [{"hypothesis": f"{o['Approach']} at step {o['Step']} ({o['Name']}) cuts cycle time and manual hours "
                                  "without raising errors",
                    "variant": "Run the change for one region, plant or distributor group for four weeks and keep the "
                               "current process elsewhere as the control group",
                    "primary_metric": "Cycle time per case, from start to finish",
                    "guardrail_metric": "Share of cases rejected, reworked or escalated",
                    "platforms": "Workflow pilot with a control group, timed from workflow tool or system logs"}
                   for o in opps[:3]]

    return {"steps": steps, "metrics": metrics, "findings": findings, "opportunities": opps,
            "experiments": experiments, "time_level": time_level}


def counts(result):
    c = {"High": 0, "Medium": 0, "Low": 0}
    for f in result["findings"]:
        c[f["severity"]] += 1
    return c


def to_excel(result, buf, client="", process_name=""):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    m, c = result["metrics"], counts(result)
    summary = [["Canyon Diagnostic Agent: Process Diagnostic", ""], ["Client", client or "Not specified"],
               ["Process", process_name or "Not specified"], ["Steps", m["steps"]], ["Owners", m["owners"]],
               ["Systems and tools", m["systems"]], ["Handoffs", m["handoffs"]], ["Approvals", m["approvals"]],
               ["Findings (High / Medium / Low)", f"{c['High']} / {c['Medium']} / {c['Low']}"]]
    if m["timed"]:
        summary += [["Active work per case (hours)", m["work_h"]], ["Waiting per case (hours)", m["wait_h"]],
                    ["Cycle time per case (days)", m["cycle_days"]], ["Share of time spent on active work %", m["flow_eff"]],
                    ["Times based on", result["time_level"]]]
    if m["cases"]:
        summary += [["Cases per month", m["cases"]], ["Manual hours per month", m["manual_h_month"]]]
    summary += [["", ""], ["Evidence levels", "1 Observed (process map or SOP), 2 Measured (system timestamps), 3 Validated (pilot)"],
                ["Note", "Automation ideas are suggestions. Validate each with a pilot on one group before rollout."]]
    sheets = {
        "Summary": pd.DataFrame(summary, columns=["Item", "Value"]),
        "Findings": pd.DataFrame([{"Severity": f["severity"], "Steps": ", ".join(map(str, f["steps"])), "Finding": f["title"],
                                   "Category": f["category"], "Evidence level": f["evidence_level"], "Evidence": f["evidence"],
                                   "Recommendation": f["recommendation"], "Effort": f["effort"]} for f in result["findings"]]),
        "Process Steps": pd.DataFrame([{"Step": s["n"], "Name": s["step"], "Owner": s["owner"], "System": s["system"],
                                        "Type": s["type"], "Manual": "Yes" if s["manual"] else "No",
                                        "Work minutes": s["work_min"], "Wait hours": s["wait_h"],
                                        "Share of cases %": s["share"], "Manual hours per month": s["hours_month"]}
                                       for s in result["steps"]]),
        "Automation": pd.DataFrame(result["opportunities"]),
        "Experiments": pd.DataFrame([{"Hypothesis": e["hypothesis"], "Variant": e["variant"],
                                      "Primary metric": e["primary_metric"], "Guardrail metric": e["guardrail_metric"],
                                      "Platforms": e["platforms"]} for e in result["experiments"]]),
    }
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            (df if len(df) else pd.DataFrame({"Info": ["Nothing reported"]})).to_excel(xw, sheet_name=name, index=False)
        head = PatternFill("solid", fgColor="0B6E69")
        fills = {"High": "F8E1DE", "Medium": "F6EAD3", "Low": "E0E9F3"}
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head
            ws.freeze_panes = "A2"
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(x.value or "")) for x in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 60)
                for x in col:
                    x.alignment = Alignment(wrap_text=True, vertical="top")
        for row in xw.book["Findings"].iter_rows(min_row=2):
            if row[0].value in fills:
                row[0].fill = PatternFill("solid", fgColor=fills[row[0].value])
