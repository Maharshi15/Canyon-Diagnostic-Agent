"""Canyon Diagnostic Agent: master data scan engine.

All counts and scores come from this file, so results are repeatable.
The AI (Artificial Intelligence) model only writes the summary text.
"""
import re
from collections import OrderedDict

import pandas as pd

UOM_OK = ["MTR", "NOS", "KG", "L", "ML", "ROLL", "BOX", "SET"]
UOM_MAP = {"M": "MTR", "MT": "MTR", "METER": "MTR", "METRE": "MTR", "EA": "NOS", "PCS": "NOS",
           "PC": "NOS", "NO": "NOS", "NUMBER": "NOS", "RL": "ROLL", "ROL": "ROLL", "KGS": "KG",
           "LTR": "L", "LITRE": "L"}
SYN = {"coupler": "coupling", "teflon": "ptfe", "carton": "box", "corrugated": "box",
       "degree": "deg", "degrees": "deg"}
NOUNS = ["end cap", "pipe", "elbow", "tee", "coupling", "reducer", "valve", "bolt", "nut",
         "cement", "primer", "tape", "box", "glue"]
MATERIALS = ["pvc", "cpvc", "upvc", "hdpe", "brass", "ptfe", "ms", "ss", "gi"]
VAGUE = re.compile(r"\b(misc|miscellaneous|test|item|general|do not use|dummy|temp)\b", re.I)
DIMS = ["Completeness", "Validity", "Consistency", "Uniqueness"]


def map_columns(columns):
    cols = [str(c) for c in columns]

    def find(pattern):
        for i, c in enumerate(cols):
            if re.search(pattern, c.strip(), re.I):
                return i
        return -1

    code = find(r"^(material|mat|item|sku|part)?\s*(code|no|number|id)$|^material$")
    desc = find(r"desc|name|text")
    uom = find(r"uom|unit|measure")
    group = find(r"group|category|class|type")
    plant = find(r"plant|site|location|branch")
    if code < 0:
        code = 0
    if desc < 0:
        desc = 1 if len(cols) > 1 else 0
    return {"code": code, "desc": desc, "uom": uom, "group": group, "plant": plant}


def norm_text(s):
    t = " " + str(s or "").lower() + " "
    t = t.replace("°", " deg ")
    t = re.sub(r'(\d)\s*"', r"\1 inch ", t)
    t = re.sub(r"(\d)\s*(inches|inch|in)\b", r"\1 inch", t)
    t = re.sub(r"(\d)\s*(mm|ml|kg|ply)\b", r"\1 \2", t)
    t = re.sub(r"\b(sch|sdr|pn)[\s-]*(\d+)", r"\1\2", t)
    t = re.sub(r"\bm(\d+)\s*x\s*(\d+)", r"m\1x\2", t)
    t = re.sub(r"[^a-z0-9.]+", " ", t)
    out = []
    for w in t.split():
        w2 = SYN.get(w, w)
        if w2 and w2 not in out:
            out.append(w2)
    return " ".join(out)


def dup_key(s):
    return " ".join(sorted(norm_text(s).split()))


def standardize(s):
    n = norm_text(s)
    if not n:
        return ""
    for noun in NOUNS:
        pat = re.compile(r"\b" + noun + r"\b")
        if pat.search(n):
            toks = pat.sub(" ", n).split()
            rest = [t for t in toks if t in MATERIALS] + [t for t in toks if t not in MATERIALS]
            return (noun + (", " + " ".join(rest) if rest else "")).upper()
    return n.upper()


def norm_group(g):
    return re.sub(r"S$", "", str(g or "").strip().upper())


def _sev(label):
    if re.search(r"duplicate records|Same material code|Unit differs|Group differs|Missing unit", label, re.I):
        return "High"
    if re.search(r"Missing|Vague|unit of measure", label, re.I):
        return "Medium"
    return "Low"


def _rec(label):
    rules = [
        (r"duplicate records", "Business owner validates each cluster, picks a surviving code, and blocks the rest from new transactions."),
        (r"Same material code", "Investigate the source systems; one code must map to one material only."),
        (r"Unit differs", "Agree one base unit per material and define conversions for the others."),
        (r"Group differs", "Fix the hierarchy mapping so one material sits in one group."),
        (r"Missing unit", "Make unit of measure mandatory at creation."),
        (r"Missing", "Fill the field and make it mandatory in the creation form."),
        (r"Non standard unit of measure", "Map to the approved unit list and restrict the dropdown."),
        (r"Vague", "Review with the owner; block or rename placeholder records."),
        (r"notation", "Adopt one convention, for example 1 INCH and 25 MM, and apply the suggested descriptions."),
        (r"case", "Apply one case convention through a naming template."),
        (r"spaces", "Trim and collapse spaces on save."),
        (r"spelled", "Merge group variants into one approved list."),
    ]
    for pat, text in rules:
        if re.search(pat, label, re.I):
            return text
    return "Review with the data owner."


def _agg_label(label):
    if "Potential duplicate" in label:
        return "Potential duplicate records"
    if label.startswith("Code "):
        return "Same material code used more than once"
    if label.startswith('Non standard unit "'):
        return "Non standard unit of measure"
    if label.startswith("Group spelled"):
        return "Group name spelled several ways"
    return re.sub(r" cluster D\d+", " a duplicate cluster", label)


def read_table(file_or_path, name=""):
    """Read CSV, TSV or Excel into a DataFrame of strings."""
    name = (name or str(file_or_path)).lower()
    if name.endswith((".xlsx", ".xls", ".xlsm")):
        return pd.read_excel(file_or_path, dtype=str).fillna("")
    sep = "\t" if name.endswith(".tsv") else ","
    return pd.read_csv(file_or_path, dtype=str, sep=sep, keep_default_na=False, skipinitialspace=False)


def scan(df):
    if df is None or len(df) == 0:
        return {"error": "The file has no records. Add a header row and at least one record."}
    df = df.fillna("").astype(str)
    cols = map_columns(df.columns)
    names = list(df.columns)

    def get(row, key):
        i = cols[key]
        return row[names[i]] if i >= 0 else ""

    recs = []
    for idx, row in df.iterrows():
        raw = get(row, "desc")
        recs.append({"line": idx + 2, "code": get(row, "code").strip(), "desc_raw": raw,
                     "desc": raw.strip(), "uom": get(row, "uom").strip(),
                     "group": get(row, "group").strip(), "plant": get(row, "plant").strip(),
                     "issues": []})

    def add(r, dim, label):
        r["issues"].append((dim, label))

    # Completeness
    for r in recs:
        if not r["desc"]:
            add(r, "Completeness", "Missing description")
        if cols["uom"] >= 0 and not r["uom"]:
            add(r, "Completeness", "Missing unit of measure")
        if cols["group"] >= 0 and not r["group"]:
            add(r, "Completeness", "Missing material group")

    # Validity
    for r in recs:
        if r["uom"]:
            u = r["uom"].upper()
            if u not in UOM_OK:
                add(r, "Validity", f'Non standard unit "{r["uom"]}"' + (f" (use {UOM_MAP[u]})" if u in UOM_MAP else ""))
        if r["desc"] and (VAGUE.search(r["desc"]) or len(norm_text(r["desc"]).split()) < 2):
            add(r, "Validity", "Vague or placeholder description")

    # Consistency
    with_desc = [r for r in recs if re.search(r"[a-z]", r["desc"], re.I)]
    caps = sum(1 for r in with_desc if r["desc"] == r["desc"].upper())
    majority_caps = caps >= len(with_desc) / 2 if with_desc else True
    for r in recs:
        d = r["desc"]
        if not d:
            continue
        if re.search(r"[a-z]", d, re.I) and (d == d.upper()) != majority_caps:
            add(r, "Consistency", "Not in upper case like most records" if majority_caps else "Upper case unlike most records")
        if re.search(r'\d\s*"|°|\d(in|mm|ml|ply)\b', d, re.I) or re.search(r"\b\d+\s?in\b", d, re.I):
            add(r, "Consistency", "Non standard unit notation in text")
        if re.search(r"\s{2,}", r["desc_raw"]) or r["desc_raw"] != r["desc_raw"].strip():
            add(r, "Consistency", "Extra spaces")
    if cols["group"] >= 0:
        variants = {}
        for r in recs:
            if r["group"]:
                variants.setdefault(norm_group(r["group"]), set()).add(r["group"])
        for r in recs:
            if r["group"] and len(variants[norm_group(r["group"])]) > 1:
                add(r, "Consistency", "Group spelled several ways (" + " / ".join(sorted(variants[norm_group(r["group"])])) + ")")

    # Uniqueness
    by_code = OrderedDict()
    for r in recs:
        if r["code"]:
            by_code.setdefault(r["code"], []).append(r)
    dup_codes = [g for g in by_code.values() if len(g) > 1]
    for g in dup_codes:
        for r in g:
            add(r, "Uniqueness", f"Code {r['code']} used {len(g)} times")

    by_key = OrderedDict()
    for r in recs:
        if r["desc"]:
            by_key.setdefault(dup_key(r["desc"]), []).append(r)
    dup_groups = []
    for g in [g for g in by_key.values() if len(g) > 1]:
        cid = "D" + str(len(dup_groups) + 1).zfill(2)
        uoms = sorted({UOM_MAP.get(r["uom"].upper(), r["uom"].upper()) for r in g if r["uom"]})
        groups = sorted({norm_group(r["group"]) for r in g if r["group"]})
        for r in g:
            add(r, "Uniqueness", f"Potential duplicate, cluster {cid}")
        dg = {"id": cid, "records": g, "std": standardize(g[0]["desc"]),
              "uom_conflict": len(uoms) > 1, "group_conflict": len(groups) > 1}
        if dg["uom_conflict"]:
            for r in g:
                add(r, "Consistency", f"Unit differs inside cluster {cid}")
        if dg["group_conflict"]:
            for r in g:
                add(r, "Consistency", f"Group differs inside cluster {cid}")
        dup_groups.append(dg)

    for r in recs:
        r["std"] = standardize(r["desc"])

    total = len(recs)
    redundant = sum(len(g["records"]) - 1 for g in dup_groups) + sum(len(g) - 1 for g in dup_codes)
    dim_scores = []
    for d in DIMS:
        hit = redundant if d == "Uniqueness" else sum(1 for r in recs if any(x[0] == d for x in r["issues"]))
        dim_scores.append({"dim": d, "affected": hit, "score": round(100 * (1 - min(hit, total) / total))})
    clean = sum(1 for r in recs if not r["issues"])
    overall = round(sum(d["score"] for d in dim_scores) / len(DIMS))

    agg = OrderedDict()
    for r in recs:
        for dim, label in r["issues"]:
            key = _agg_label(label)
            f = agg.setdefault(key, {"dim": dim, "label": key, "rows": []})
            if r not in f["rows"]:
                f["rows"].append(r)
    rank = {"High": 0, "Medium": 1, "Low": 2}
    findings = []
    for f in agg.values():
        findings.append({**f, "count": len(f["rows"]), "severity": _sev(f["label"]), "rec": _rec(f["label"])})
    findings.sort(key=lambda f: (rank[f["severity"]], -f["count"]))

    return {"cols": cols, "columns": names, "recs": recs, "total": total, "clean": clean,
            "overall": overall, "dim_scores": dim_scores, "dup_groups": dup_groups,
            "dup_codes": dup_codes, "findings": findings}


def to_excel(result, path_or_buffer, client_name=""):
    """Write the four sheet diagnostic workbook."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    summary = [["Canyon Diagnostic Agent: Master Data Diagnostic", ""],
               ["Client", client_name or "Not specified"],
               ["Records reviewed", result["total"]],
               ["Data quality score (out of 100)", result["overall"]]]
    summary += [[f"{d['dim']} score", d["score"]] for d in result["dim_scores"]]
    summary += [["Potential duplicate clusters (to be validated)", len(result["dup_groups"])],
                ["Records with at least one issue", result["total"] - result["clean"]],
                ["", ""],
                ["Note", "The agent flags potential issues. Business owners approve every correction before master data is changed."]]
    findings = pd.DataFrame([{"Severity": f["severity"], "Finding": f["label"], "Dimension": f["dim"],
                              "Records affected": f["count"], "Evidence level": "2 Measured (data extract)",
                              "Example records": "; ".join(f"{r['code']} {r['desc']}" for r in f["rows"][:5]),
                              "Recommendation": f["rec"]} for f in result["findings"]])
    clusters = pd.DataFrame([{"Cluster": g["id"], "Suggested standard description": g["std"],
                              "Code": r["code"], "Description": r["desc"], "Unit": r["uom"],
                              "Group": r["group"], "Plant": r["plant"],
                              "Unit conflict": "Yes" if g["uom_conflict"] else "",
                              "Group conflict": "Yes" if g["group_conflict"] else "",
                              "Keep (business owner to mark)": ""}
                             for g in result["dup_groups"] for r in g["records"]])
    records = pd.DataFrame([{"Row": r["line"], "Code": r["code"], "Description": r["desc"], "Unit": r["uom"],
                             "Group": r["group"], "Plant": r["plant"], "Suggested standard description": r["std"],
                             "Issues": "; ".join(OrderedDict.fromkeys(x[1] for x in r["issues"])) or "Clean"}
                            for r in result["recs"]])

    with pd.ExcelWriter(path_or_buffer, engine="openpyxl") as xw:
        pd.DataFrame(summary, columns=["Item", "Value"]).to_excel(xw, sheet_name="Summary", index=False)
        findings.to_excel(xw, sheet_name="Findings", index=False)
        clusters.to_excel(xw, sheet_name="Duplicate Clusters", index=False)
        records.to_excel(xw, sheet_name="All Records", index=False)
        head = PatternFill("solid", fgColor="0B6E69")
        sev_fill = {"High": "F8E1DE", "Medium": "F6EAD3", "Low": "E0E9F3"}
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head
            ws.freeze_panes = "A2"
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 70)
                for c in col:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        ws = xw.book["Findings"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value in sev_fill:
                row[0].fill = PatternFill("solid", fgColor=sev_fill[row[0].value])
