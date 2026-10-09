"""Canyon Diagnostic Agent: client Diagnostic Report (Word).

Combines whichever modules were run in this session into one evidence graded report.
Every number comes from the module results. Client facing text follows Canyon house style:
plain English, no dash characters, abbreviations spelled out the first time.
"""
import io
import json
import re
from datetime import date

import app_engine as AE
import process_engine as PE

TEAL, INK, MUTED = "0B6E69", "132228", "5B6B70"
SEV_FILL = {"High": "F8E1DE", "Medium": "F6EAD3", "Low": "E0E9F3"}
SEV_RANK = {"High": 0, "Medium": 1, "Low": 2}
NEXT_STEP = "Expert consultation with Canyon to agree priorities and a transformation roadmap."

# Master data findings carry no effort in the scan, so the report assigns one here
MD_EFFORT = [(r"duplicate", "M"), (r"Same material code", "M"), (r"Unit differs|Group differs", "M"),
             (r"Missing|Vague|Non standard", "S"), (r".", "S")]
MD_IMPACT = {"High": "Wrong transactions or wrong reporting", "Medium": "Rework and slower teams",
             "Low": "Hygiene and naming convention"}


def clean(text):
    """House style guard: no dash characters in client facing text."""
    t = str(text or "")
    t = re.sub(r"\s*[\u2013\u2014]\s*", ", ", t)
    t = re.sub(r"\s+-\s+", ", ", t)
    t = re.sub(r"(?<=[A-Za-z])-(?=[A-Za-z])", " ", t)
    return t.replace("\u2192", "to")


# Collect findings from each module into one shape

def collect(md=None, app=None, proc=None):
    """Each argument is the module's session data, or None. Returns a dict the Word writer uses."""
    mods, findings, experiments, quick, roadmap = [], [], [], [], {"Fix": [], "Simplify": [], "Automate": [], "Transform": []}
    illustrative = False

    if md:
        res = md["res"]
        illustrative |= "sample" in md.get("source", "").lower()
        dims = ", ".join(f"{d['dim']} {d['score']}" for d in res["dim_scores"])
        mods.append({"name": "Master Data Diagnostic", "reviewed": f"{res['total']} records from {md.get('source', 'a data extract')}",
                     "level": "2 Measured (data extract)",
                     "headline": f"Data quality score {res['overall']} out of 100",
                     "detail": f"{dims}. {len(res['dup_groups'])} potential duplicate clusters to validate. "
                               f"{res['total'] - res['clean']} of {res['total']} records have at least one issue.",
                     "gaps": ["Business owner validation of each duplicate cluster",
                              "Customer, distributor and vendor masters, if not yet shared"]})
        for f in res["findings"]:
            effort = next(e for p, e in MD_EFFORT if re.search(p, f["label"], re.I))
            ex = "; ".join(r["desc"] for r in f["rows"][:3] if r["desc"])
            label = f["label"].replace("inside a duplicate cluster", "inside a potential duplicate cluster")
            label += ", to be validated by the business owner" if label.startswith("Potential duplicate") else ""
            findings.append({"module": "MD", "finding": label, "severity": f["severity"], "level": "2 Measured",
                             "evidence": f"{f['count']} of {res['total']} records." + (f" Examples: {ex}." if ex else ""),
                             "rec": f["rec"], "effort": effort, "impact": MD_IMPACT[f["severity"]]})
        roadmap["Fix"].append(f"Cleanse the master data: fill missing attributes, apply standard descriptions and units, "
                              f"and retire confirmed duplicates (score today {res['overall']} out of 100).")
        roadmap["Transform"].append("Data governance: creation templates, mandatory fields and business owner approval "
                                    "so the master stays clean after cleansing.")

    if app:
        funnel = AE.funnel_analysis(app.get("funnel"))
        usage = AE.page_usage_analysis(app.get("usage"))
        res = AE.attach_measured(AE.normalize(json.loads(json.dumps(app["result"])), len(app["files"])), funnel)
        illustrative |= "sample" in app.get("source", "").lower()
        c, j = AE.counts(res), res["journey"]
        detail = f"{c['High']} High, {c['Medium']} Medium, {c['Low']} Low findings. Journey from {j.get('current_steps')} to {j.get('proposed_steps')} steps."
        if funnel:
            detail += f" Measured completion {funnel['completion']}% ({funnel['end']} of {funnel['start']} users)."
        if usage:
            detail += f" {len(usage['low'])} pages with under {usage['threshold']}% of all views."
        gaps = []
        if not funnel:
            gaps.append("Funnel export: without it, drop off is observed, not measured")
        if not usage:
            gaps.append("Page usage export: without it, unused pages cannot be confirmed")
        gaps.append("A/B (split) test results to validate the proposed journey")
        name = app.get("app_name") or "the application"
        mods.append({"name": "Application Diagnostic", "reviewed": f"{len(app['files'])} screens of {name}",
                     "level": "2 Measured (analytics)" if funnel or usage else "1 Observed (screenshots)",
                     "headline": f"Journey can go from {j.get('current_steps')} to {j.get('proposed_steps')} steps",
                     "detail": detail, "gaps": gaps})
        for f in res["findings"]:
            ev = f"Screen {f['screen']}: {f['evidence']}" + (f" Measured: {f['measured']}." if f["measured"] else "")
            findings.append({"module": "AP", "finding": f["title"], "severity": f["severity"], "level": f["evidence_level"],
                             "evidence": ev, "rec": f["recommendation"], "effort": f["effort"],
                             "impact": "Users blocked or lost" if f["severity"] == "High" else
                                       "Slower users and rework" if f["severity"] == "Medium" else "Polish"})
        if usage and usage["low"]:
            findings.append({"module": "AP", "finding": "Pages with very low usage", "severity": "Low", "level": "2 Measured",
                             "evidence": "; ".join(f"{r['Page']} ({r['Share of all views %']}% of views)" for r in usage["low"]),
                             "rec": "Review with the product owner: merge, hide or retire these pages.", "effort": "S",
                             "impact": "Simpler navigation and less to maintain"})
        for e in res["experiments"]:
            experiments.append({"module": "Application", **{k: e.get(k, "") for k in
                                ("hypothesis", "variant", "primary_metric", "guardrail_metric", "platforms")}})
        roadmap["Simplify"].append(f"Redesign the {app.get('app_name') or 'application'} journey from "
                                   f"{j.get('current_steps')} to {j.get('proposed_steps')} steps, validated with an A/B test.")

    if proc:
        res = proc["result"]
        m = res["metrics"]
        illustrative |= "sample" in proc.get("source", "").lower()
        c = PE.counts(res)
        detail = f"{m['steps']} steps, {m['owners']} owners, {m['handoffs']} handoffs, {m['approvals']} approvals."
        if m["timed"]:
            detail += f" Cycle time {m['cycle_days']:g} days per case, {m['flow_eff']:g}% of it active work."
        if m["cases"]:
            detail += f" {m['manual_h_month']:,.0f} manual hours per month for {m['cases']:,} cases."
        pname = proc.get("process_name") or "the process"
        headline = (f"Cycle time {m['cycle_days']:g} days, {m['flow_eff']:g}% active work" if m["timed"]
                    else f"{c['High']} High and {c['Medium']} Medium findings")
        gaps = []
        if "Measured" not in res["time_level"]:
            gaps.append("System timestamps or logs: times today are estimates")
        if not m["cases"]:
            gaps.append("Monthly case volume: needed to size manual hours")
        gaps.append("A pilot on one group to validate the automation")
        mods.append({"name": "Process Diagnostic", "reviewed": f"{m['steps']} steps of {pname}",
                     "level": res["time_level"] + (" (times)" if m["timed"] else ""), "headline": headline,
                     "detail": detail, "gaps": gaps})
        for f in res["findings"]:
            findings.append({"module": "PR", "finding": f["title"], "severity": f["severity"], "level": f["evidence_level"],
                             "evidence": f["evidence"], "rec": f["recommendation"], "effort": f["effort"],
                             "impact": "Delays, errors and manual effort" if f["severity"] == "High" else "Slower cycle and rework"})
        for e in res["experiments"]:
            experiments.append({"module": "Process", **e})
        top = [o for o in res["opportunities"] if o["Impact"] != "Low"][:3]
        if top:
            roadmap["Automate"].append(f"Automate {pname}: " + "; ".join(f"step {o['Step']} {o['Name']} ({o['Approach'][0].lower() + o['Approach'][1:]})" for o in top) + ".")
        roadmap["Simplify"].append(f"Agree approval limits and remove handoffs in {pname}.")

    findings.sort(key=lambda f: SEV_RANK[f["severity"]])
    n = {}
    for f in findings:
        n[f["module"]] = n.get(f["module"], 0) + 1
        f["id"] = f"{f['module']}{n[f['module']]}"
    quick = [f for f in findings if f["effort"] == "S" and f["severity"] in ("High", "Medium")][:6]
    if roadmap["Automate"] or roadmap["Fix"]:
        roadmap["Transform"].append("Connect the systems involved on the Canyon platform, so business teams see one "
                                    "reconciled view of operations and can ask questions of the data directly.")
    return {"modules": mods, "findings": findings, "experiments": experiments, "quick": quick,
            "roadmap": roadmap, "illustrative": illustrative}


def exec_summary(data, client):
    mods, f = data["modules"], data["findings"]
    c = {s: sum(1 for x in f if x["severity"] == s) for s in SEV_RANK}
    names = [m["name"].replace(" Diagnostic", "").lower() for m in mods]
    scope = names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1]
    lines = [f"Canyon Data Labs reviewed the {scope} of {client or 'the client'}. "
             f"We found {len(f)} findings: {c['High']} High, {c['Medium']} Medium and {c['Low']} Low."]
    lines += [f"{m['name']}: {m['headline']}." for m in mods]
    # Top issues: the most severe finding of each module first, then the next most severe overall
    top = []
    for code in ("MD", "AP", "PR"):
        top += [x for x in f if x["module"] == code][:1]
    top = sorted(top, key=lambda x: SEV_RANK[x["severity"]])[:3]
    top += [x for x in f if x not in top][:3 - len(top)]
    if top:
        lines.append("Top issues: " + "; ".join(x["finding"] for x in top) + ".")
    lines.append(f"Recommended next step: {NEXT_STEP[0].lower() + NEXT_STEP[1:]}")
    return [clean(x) for x in lines]


# Word writer

def _shade(cell, hex_fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tc.append(sh)


def _table(doc, header, rows, widths=None, sev_col=None, size=9):
    from docx.shared import Cm, Pt, RGBColor
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(size), RGBColor(255, 255, 255)
        _shade(cell, TEAL)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(clean(v)).font.size = Pt(size)
            if sev_col is not None and i == sev_col and v in SEV_FILL:
                _shade(cells[i], SEV_FILL[v])
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def build_docx(data, client="", prepared_by="", contact="", notes="", teaser=False):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = docx.Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.color.rgb = "Calibri", Pt(size), RGBColor.from_string(TEAL)

    def para(text, bold=False, size=None, color=None, italic=False, align=None):
        p = doc.add_paragraph()
        r = p.add_run(clean(text))
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if align:
            p.alignment = align
        return p

    def bullet(text):
        doc.add_paragraph(clean(text), style="List Bullet")

    title = "Canyon Diagnostic Summary" if teaser else "Canyon Diagnostic Report"
    para("CANYON DATA LABS", bold=True, size=9, color=TEAL)
    para(title, bold=True, size=22, color=INK)
    para(f"Prepared for {client or 'the client'}  ·  {date.today():%d %B %Y}", size=11, color=MUTED)
    if prepared_by:
        para(f"Prepared by {prepared_by}" + (f", {contact}" if contact else ""), size=10, color=MUTED)
    para("Before you transform your business systems, diagnose them.", italic=True, size=10, color=TEAL)
    if data["illustrative"]:
        para("Illustrative: this report uses sample data, not client data.", bold=True, size=10, color="B3372B")
    para(f"Confidential. Prepared for {client or 'the client'} only.", size=8.5, color=MUTED)

    doc.add_heading("1. Executive summary", level=1)
    for line in exec_summary(data, client):
        bullet(line)
    if notes.strip():
        para(notes.strip())

    if not teaser:
        doc.add_heading("2. Scope and inputs", level=1)
        _table(doc, ["Module", "What we reviewed", "Evidence level", "Not received yet, and what it limits"],
               [[m["name"], m["reviewed"], m["level"], "; ".join(m["gaps"])] for m in data["modules"]],
               widths=[3.4, 4.2, 3.4, 6])
        para("Evidence levels: 1 Observed (seen in screenshots or documents, can suggest but not prove), "
             "2 Measured (backed by data extracts, analytics or logs), 3 Validated (confirmed by an experiment or pilot).",
             size=9, color=MUTED)

    doc.add_heading(("2" if teaser else "3") + ". Scorecard", level=1)
    _table(doc, ["Module", "Headline", "Detail"], [[m["name"], m["headline"], m["detail"]] for m in data["modules"]],
           widths=[3.6, 4.8, 8.6])

    doc.add_heading(("3" if teaser else "4") + ". Findings" + (": top five" if teaser else ""), level=1)
    rows = data["findings"][:5] if teaser else data["findings"]
    _table(doc, ["ID", "Severity", "Level", "Finding and evidence", "Recommendation", "Effort"],
           [[f["id"], f["severity"], f["level"].split()[0], f"{f['finding']}. {f['evidence']}", f["rec"], f["effort"]]
            for f in rows], widths=[1.2, 1.7, 1.2, 6.6, 5.1, 1.2], sev_col=1, size=8.5)
    para("Severity: High means wrong transactions, wrong reporting or blocked users. Medium means rework or slower users. "
         "Low means hygiene and convention. Effort: S under 30 days, M one to three months, L longer.", size=9, color=MUTED)

    if not teaser:
        doc.add_heading("5. Quick wins (under 30 days)", level=1)
        if data["quick"]:
            for f in data["quick"]:
                bullet(f"{f['id']}: {f['rec']}")
        else:
            para("No low effort fixes stood out. Start with the highest severity findings above.")

        doc.add_heading("6. Experiments to validate", level=1)
        if data["experiments"]:
            _table(doc, ["Module", "Hypothesis", "Design", "Primary metric", "Guardrail"],
                   [[e["module"], e["hypothesis"], e["variant"], e["primary_metric"], e["guardrail_metric"]]
                    for e in data["experiments"]], widths=[2, 5, 4.6, 2.8, 2.6], size=8.5)
        else:
            para("Experiments apply to Application and Process findings. None were in scope for this diagnostic.")

        doc.add_heading("7. Roadmap", level=1)
        offer = {"Fix": "Data engineering and cleansing", "Simplify": "Application modernization and process redesign",
                 "Automate": "Workflow automation and AI (Artificial Intelligence) agents",
                 "Transform": "Data governance and the Canyon platform"}
        _table(doc, ["Stage", "What it means here", "How Canyon helps"],
               [[k, " ".join(v) if v else "Not in scope for this diagnostic", offer[k]] for k, v in data["roadmap"].items()],
               widths=[2.4, 9.6, 5])

    doc.add_heading(("4" if teaser else "8") + ". Next step", level=1)
    para(NEXT_STEP, bold=True)
    para("Canyon works with the systems a business already runs on, no need to replace anything. Our flagship "
         "deployment reconciles data across more than 2,600 distributors and 300,000 retailers for a major Indian manufacturer.")
    if prepared_by or contact:
        para(f"Contact: {', '.join(x for x in (prepared_by, contact) if x)}", color=MUTED)

    foot = sec.footer.paragraphs[0]
    foot.text = f"Canyon Data Labs, Ahmedabad  ·  {title} for {client or 'the client'}  ·  Confidential"
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in foot.runs:
        r.font.size, r.font.color.rgb = Pt(8), RGBColor.from_string(MUTED)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
